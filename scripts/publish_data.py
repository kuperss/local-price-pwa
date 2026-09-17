"""V36 cache → approved SKU list → encrypted JSON → authenticated device downloads.

No Excel prices are imported: the workbook establishes scope only. No plaintext products
or keys are written into the public static directory. D1's current pointer advances last.
"""
import argparse
import base64
import gzip
import hashlib
import json
import os
import tempfile
import unicodedata
import uuid
from datetime import datetime, timezone
from pathlib import Path
import openpyxl
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cloud import Cloud, ROOT, deployment
import catalog_sync

BASE=ROOT.parent.parent/'價格查詢工具'/'價格查詢'
COST_SCHEMA=[sql.strip() for sql in (ROOT/'worker/schema.sql').read_text(encoding='utf-8').split(';')
             if sql.strip().startswith(('CREATE TABLE IF NOT EXISTS cost_bundles ', 'CREATE TABLE IF NOT EXISTS cost_chunks('))]
def clean(code):
    return str(code).strip().replace('"','').replace('\\','').replace('/','').upper()
def canonical(code):
    return unicodedata.normalize('NFKC',str(code)).strip().upper()
def utc():return datetime.now(timezone.utc).isoformat()
def read_seed(path):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    try:
        ws=wb['合併結果'] if '合併結果' in wb.sheetnames else wb.worksheets[0]
        rows=ws.iter_rows(values_only=True);headers=list(next(rows));idx=headers.index('料號')
        codes=set();invalid=0
        for row in rows:
            value=row[idx]
            if not value:continue
            code=canonical(value)
            if code.startswith(('查無','已中斷')):invalid+=1;continue
            codes.add(code)
        return sorted(codes),invalid
    finally:wb.close()

def product(row,specs):
    mapping={'IMA01':'型號','IMA02':'中文品名','IMA130':'售架','IMA140':'停產','TA_IMA107':'售轉料號','IMA021':'包裝規格(小/中/大)','A':'底價','B':'量價','C':'開盤價','IMAUD01':'備註'}
    out={mapping.get(k,k):str(v).strip() for k,v in row.items() if v is not None and str(v).strip() and k not in ['搭贈1','搭贈2','搭贈3']}
    bonus=[str(row.get(k) or '').strip() for k in ['搭贈1','搭贈2','搭贈3']]
    if any(bonus):out['搭贈']=' / '.join(v for v in bonus if v)
    for prefix,fields in [('單入尺寸(cm)',['單入長(cm)','單入寬(cm)','單入高(cm)']),('外箱尺寸(cm)',['外箱長度(cm)','外箱寬度(cm)','外箱高度(cm)'])]:
        if all(out.get(f) for f in fields):
            out[prefix]=' x '.join(out.pop(f) for f in fields)
    for k,v in specs.items():
        if v is not None and str(v).strip():out['規格_'+k]=str(v).strip()
    return out

def build(payload,codes):
    master=payload['master']; specs=payload.get('specs',{})
    raw_index={canonical(row.get('IMA01',key)):(key,row) for key,row in master.items()}
    cleaned_index={clean(key):(key,row) for key,row in master.items()}
    rows=[];missing=[];resolved=[];seen=set()
    for code in sorted(codes):
        hit=raw_index.get(canonical(code)) or cleaned_index.get(clean(code))
        if not hit:missing.append(code);continue
        key,row=hit;resolved.append(code)
        if key in seen:continue
        seen.add(key);rows.append(product(row,specs.get(key,{})))
    if not rows:raise ValueError('No approved products resolved; previous bundle retained')
    return rows,resolved,missing

def atomic(path,data):
    path.parent.mkdir(parents=True,exist_ok=True)
    fd,tmp=tempfile.mkstemp(prefix=path.name,dir=path.parent)
    try:
        with os.fdopen(fd,'w',encoding='utf-8') as f:json.dump(data,f,ensure_ascii=False)
        os.replace(tmp,path)
    finally:
        if os.path.exists(tmp):os.unlink(tmp)


def split_costs(rows):
    prices=[];costs=[]
    for row in rows:
        price={};cost={'型號':str(row.get('型號',''))}
        for k,v in row.items():
            label=unicodedata.normalize('NFKC',k)
            if '成本' in label or 'cost' in label.lower():
                cost[k]=v
            else:price[k]=v
        prices.append(price)
        if len(cost)>1:costs.append(cost)
    return prices,costs

def publish_bundle(cloud,db,rows,fetched,resolved,missing,invalid=0,catalog_plan=None):
    # The scheduled job uses this checkout: create additive tables before any upload,
    # even when the new Worker has not been deployed yet. General prices remain usable.
    for sql in COST_SCHEMA:cloud.query(db,sql)
    stamp=utc()
    original=json.dumps(rows,ensure_ascii=False,separators=(',',':')).encode()
    prices,cost_rows=split_costs(rows)
    plain=json.dumps(prices,ensure_ascii=False,separators=(',',':')).encode()
    # Include freshness in version so users can verify the latest daily check even if prices are unchanged.
    revision=('\n'+catalog_plan['revision']).encode() if catalog_plan else b''
    version=hashlib.sha256(b'split-costs-v1\n'+fetched.encode()+b'\n'+original+revision).hexdigest()[:32]
    product_count=sum(not row.get('_replacements') for row in rows)
    previous=cloud.query(db,"SELECT value FROM settings WHERE key='current_bundle'")
    prior_version=previous[0]['value'] if previous else ''
    fingerprint=version
    prior_report=cloud.query(db,'SELECT report FROM catalog_publications WHERE version=?',[prior_version]) if catalog_plan else []
    unchanged=(json.loads(prior_report[0]['report']).get('fingerprint')==fingerprint) if prior_report else prior_version==version
    if unchanged:
        print(f'PRICE_PWA_OK unchanged version={prior_version} products={product_count}');return
    # Independent staging IDs prevent two concurrent uploads from mixing encryption chunks/keys.
    if catalog_plan:version=fingerprint[:20]+uuid.uuid4().hex[:12]
    cost_key=AESGCM.generate_key(bit_length=256);cost_iv=os.urandom(12)
    cost_plain=json.dumps(cost_rows,ensure_ascii=False,separators=(',',':')).encode()
    cost_cipher=AESGCM(cost_key).encrypt(cost_iv,cost_plain,('costs:'+version).encode())
    cost_encoded=base64.b64encode(cost_cipher).decode()
    cost_chunks=[cost_encoded[i:i+80000] for i in range(0,len(cost_encoded),80000)]
    key=AESGCM.generate_key(bit_length=256);iv=os.urandom(12);cipher=AESGCM(key).encrypt(iv,plain,version.encode())
    encoded=base64.b64encode(cipher).decode();chunks=[encoded[i:i+80000] for i in range(0,len(encoded),80000)]
    bundle={'format':'price-pwa-aes-gcm-v1','securityFormat':'split-costs-v1','version':version,'fetchedAt':fetched,'count':product_count,'iv':base64.b64encode(iv).decode(),'hash':hashlib.sha256(cipher).hexdigest(),'cipher':encoded}
    # This local .json intentionally contains no decryption key; keys require device authorization.
    private=ROOT/'data-private';atomic(private/'price-data.json',bundle)
    atomic(private/'publish-report.json',{'fetched_at':fetched,'products':product_count,'aliases':len(rows)-product_count,'approved':len(resolved)+len(missing),'invalid_seed_rows':invalid,'missing_skus':missing,'blocked':catalog_plan['blocked'] if catalog_plan else []})
    for i,chunk in enumerate(chunks):cloud.query(db,'INSERT OR REPLACE INTO bundle_chunks VALUES(?,?,?)',[version,i,chunk])
    cloud.query(db,'INSERT OR REPLACE INTO bundles(version,fetched_at,published_at,product_count,key_b64,iv_b64,content_hash,chunks) VALUES(?,?,?,?,?,?,?,?)',[version,fetched,stamp,product_count,base64.b64encode(key).decode(),bundle['iv'],bundle['hash'],len(chunks)])
    for i,chunk in enumerate(cost_chunks):
        cloud.query(db,'INSERT OR REPLACE INTO cost_chunks VALUES(?,?,?)',[version,i,chunk])
    cloud.query(db,'INSERT OR REPLACE INTO cost_bundles VALUES(?,?,?,?,?,?)',
                [version,base64.b64encode(cost_key).decode(),base64.b64encode(cost_iv).decode(),hashlib.sha256(cost_cipher).hexdigest(),len(cost_chunks),len(cost_rows)])
    # Publish last: interrupted upload cannot replace the last complete version.
    if catalog_plan:
        for start in range(0,len(resolved),200):
            cloud.query(db,'INSERT OR IGNORE INTO catalog_members(version,sku) SELECT ?,value FROM json_each(?)',[version,json.dumps(resolved[start:start+200])])
        report={'blocked':catalog_plan['blocked'],'missing':missing,'aliases':len(catalog_plan['aliases']),'fingerprint':fingerprint}
        cloud.query(db,'INSERT OR REPLACE INTO catalog_publications VALUES(?,?,?)',[version,catalog_plan['revision'],json.dumps(report,ensure_ascii=False)])
        published=cloud.query(db,"""INSERT INTO settings(key,value) SELECT 'current_bundle',?
          WHERE (SELECT value FROM settings WHERE key='catalog_revision')=?
          AND COALESCE((SELECT value FROM settings WHERE key='current_bundle'),'')=?
          ON CONFLICT(key) DO UPDATE SET value=excluded.value RETURNING value""",[version,catalog_plan['revision'],prior_version])
        if not published:raise ValueError('Catalog changed during upload; previous bundle retained. Retry next publish.')
    else:
        cloud.query(db,"INSERT INTO settings VALUES('current_bundle',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",[version])
    # Legacy request-screen hints follow the successful pointer, never an interrupted upload.
    for status,items in [('included',resolved),('missing',missing)]:
        for start in range(0,len(items),50):
            batch=items[start:start+50];cloud.query(db,'UPDATE catalog SET resolution=?,resolved_at=?,included_version=? WHERE sku IN ('+','.join('?' for _ in batch)+')',[status,stamp,version if status=='included' else None,*batch])
    print(f'PRICE_PWA_OK version={version} products={product_count} aliases={len(rows)-product_count} approved={len(resolved)+len(missing)} missing={len(missing)} fetched_at={fetched}')


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--seed',type=Path,default=BASE/'0903.xlsx');ap.add_argument('--cache',type=Path,default=BASE/'SEVICache.json.gz');ap.add_argument('--check-only',action='store_true');args=ap.parse_args()
    seeds,invalid=read_seed(args.seed)
    with gzip.open(args.cache,'rt',encoding='utf-8') as f:payload=json.load(f)
    fetched=payload['fetched_at']
    if not args.check_only and (datetime.now()-datetime.fromisoformat(fetched)).total_seconds()>24*3600:
        raise ValueError('V36 cache older than 24 hours; refusing to publish stale prices')
    if args.check_only:
        rows,resolved,missing=build(payload,seeds)
        print(json.dumps({'seed_unique':len(seeds),'invalid_source_rows':invalid,'products':len(rows),'missing':len(missing),'missing_skus':missing,'fetched_at':fetched},ensure_ascii=False));return
    cfg=deployment();cloud=Cloud();db=cfg['database_id'];stamp=utc()
    # Bootstrap once; subsequent runs use the approved catalog as the source of scope.
    seeded=cloud.query(db,"SELECT value FROM settings WHERE key='seed_initialized'")
    if not seeded:
        for start in range(0,len(seeds),25):
            batch=seeds[start:start+25];cloud.query(db,'INSERT OR IGNORE INTO catalog(sku,source,approved_at) VALUES '+','.join("(?,'seed',?)" for _ in batch),[v for code in batch for v in (code,stamp)])
        cloud.query(db,"INSERT OR IGNORE INTO settings VALUES('seed_initialized',?)",[stamp])
    catalog_sync.sync_metadata(cloud,db,payload)
    publication=catalog_sync.plan(cloud,db,payload)
    rows,resolved,missing=build(payload,publication['codes'])
    # Never silently drop a formerly published product because a source row vanished.
    live={r['sku'] for r in cloud.query(db,"SELECT sku FROM catalog_members WHERE version=(SELECT value FROM settings WHERE key='current_bundle')")}
    unsafe=set(missing)&(live|{r['sku'] for r in publication['blocked']})
    if unsafe:raise ValueError('Previously published/retained products missing from cache; previous bundle retained: '+', '.join(sorted(unsafe)))
    rows.extend(publication['aliases'])
    publish_bundle(cloud,db,rows,fetched,resolved,missing,invalid,publication)

if __name__=='__main__':
    try:main()
    except Exception as exc:
        print(f'PRICE_PWA_ERROR {type(exc).__name__}: {exc}');raise SystemExit(1)
